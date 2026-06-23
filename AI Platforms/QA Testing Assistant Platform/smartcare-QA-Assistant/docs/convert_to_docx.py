"""
Convert markdown files to formatted Word (.docx) documents.

Usage examples:
    python docs/convert_to_docx.py
    python docs/convert_to_docx.py --input docs/My_Document.md
    python docs/convert_to_docx.py --input docs/My_Document.md --output docs/My_Document.docx
"""
import argparse
import subprocess
import sys

# Auto-install python-docx if not present
try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("Installing python-docx ...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

from pathlib import Path
import re

# Use current working directory or script directory
SCRIPT_DIR = Path(__file__).parent if "__file__" in globals() else Path.cwd()
DEFAULT_MD_PATH = SCRIPT_DIR / "SmartCare_QA_AI_Assistant_Architecture.md"
DEFAULT_OUT_PATH = SCRIPT_DIR / "SmartCare_QA_AI_Assistant_Architecture.docx"

BRAND_PURPLE = RGBColor(0x4E, 0x0F, 0x78)
BRAND_MAROON = RGBColor(0x6B, 0x00, 0x42)
DARK_TEXT    = RGBColor(0x1D, 0x27, 0x38)

def set_heading_style(run, level):
    if level == 1:
        run.font.size   = Pt(22)
        run.font.bold   = True
        run.font.color.rgb = BRAND_PURPLE
    elif level == 2:
        run.font.size   = Pt(16)
        run.font.bold   = True
        run.font.color.rgb = BRAND_PURPLE
    elif level == 3:
        run.font.size   = Pt(13)
        run.font.bold   = True
        run.font.color.rgb = BRAND_MAROON

def add_table_from_lines(doc, lines):
    """Parse markdown table lines into a Word table."""
    rows = [l for l in lines if l.strip().startswith("|") and not re.match(r"^\|[-| :]+\|$", l.strip())]
    if not rows:
        return
    parsed = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    col_count = max(len(r) for r in parsed)

    table = doc.add_table(rows=len(parsed), cols=col_count)
    table.style = "Table Grid"

    for r_idx, row_data in enumerate(parsed):
        for c_idx, cell_text in enumerate(row_data):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = cell_text
            run = cell.paragraphs[0].runs[0] if cell.paragraphs[0].runs else cell.paragraphs[0].add_run(cell_text)
            run.font.size = Pt(10)
            if r_idx == 0:
                run.bold = True
                run.font.color.rgb = BRAND_PURPLE

    doc.add_paragraph()  # spacing after table

def build_docx(md_text: str):
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.left_margin   = Inches(1)
        section.right_margin  = Inches(1)
        section.top_margin    = Inches(1)
        section.bottom_margin = Inches(1)

    lines = md_text.splitlines()
    i = 0
    table_buffer = []
    in_table = False
    in_code  = False
    code_buffer = []

    while i < len(lines):
        line = lines[i]

        # Code block
        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_buffer = []
            else:
                in_code = False
                code_para = doc.add_paragraph()
                code_para.style = "Normal"
                run = code_para.add_run("\n".join(code_buffer))
                run.font.name = "Courier New"
                run.font.size = Pt(9)
                code_para.paragraph_format.left_indent = Inches(0.4)
            i += 1
            continue

        if in_code:
            code_buffer.append(line)
            i += 1
            continue

        # Table detection
        is_table_row = line.strip().startswith("|")
        if is_table_row:
            table_buffer.append(line)
            i += 1
            continue
        elif table_buffer:
            add_table_from_lines(doc, table_buffer)
            table_buffer = []

        # Headings
        h1 = re.match(r"^# (.+)$", line)
        h2 = re.match(r"^## (.+)$", line)
        h3 = re.match(r"^### (.+)$", line)

        if h1:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(h1.group(1))
            set_heading_style(run, 1)
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after  = Pt(6)
        elif h2:
            p = doc.add_paragraph()
            run = p.add_run(h2.group(1))
            set_heading_style(run, 2)
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after  = Pt(4)
        elif h3:
            p = doc.add_paragraph()
            run = p.add_run(h3.group(1))
            set_heading_style(run, 3)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after  = Pt(2)
        elif line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(line[2:].strip())
            run.font.size = Pt(11)
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            run = p.add_run(re.sub(r"^\d+\. ", "", line).strip())
            run.font.size = Pt(11)
        elif line.strip() == "---":
            doc.add_paragraph("-" * 80)
        elif line.strip() == "":
            doc.add_paragraph()
        else:
            p = doc.add_paragraph()
            run = p.add_run(line.strip())
            run.font.size = Pt(11)
            run.font.color.rgb = DARK_TEXT

        i += 1

    # flush remaining table
    if table_buffer:
        add_table_from_lines(doc, table_buffer)

    return doc


def resolve_paths() -> tuple[Path, Path]:
    parser = argparse.ArgumentParser(description="Convert markdown to Word (.docx).")
    parser.add_argument("--input", dest="input_path", default=str(DEFAULT_MD_PATH), help="Path to input markdown file")
    parser.add_argument("--output", dest="output_path", default="", help="Path to output docx file")
    args = parser.parse_args()

    md_path = Path(args.input_path).resolve()
    if args.output_path:
        out_path = Path(args.output_path).resolve()
    else:
        out_path = md_path.with_suffix(".docx")

    return md_path, out_path

if __name__ == "__main__":
    MD_PATH, OUT_PATH = resolve_paths()
    print(f"Reading  : {MD_PATH}")
    md_text = MD_PATH.read_text(encoding="utf-8")
    # Replace UTF-8 em-dashes with ASCII alternatives to avoid encoding issues in Word
    md_text = md_text.replace("—", " - ").replace("–", "-")
    doc = build_docx(md_text)
    doc.save(OUT_PATH)
    print(f"Saved    : {OUT_PATH}")
